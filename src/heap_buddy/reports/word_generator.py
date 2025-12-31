"""
Word Document Report Generator for Heap Analytics reports
"""

import os
from typing import Dict, List, Any, Optional
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class WordReportGenerator:
    """
    Generates professional Word documents from Heap Analytics data
    """

    # Color scheme (RGB tuples)
    COLORS = {
        "primary": (79, 70, 229),      # Indigo
        "secondary": (16, 185, 129),   # Emerald
        "accent": (245, 158, 11),      # Amber
        "danger": (239, 68, 68),       # Red
        "text": (31, 41, 55),          # Dark gray
        "text_light": (107, 114, 128), # Medium gray
    }

    def __init__(
        self,
        output_path: str,
        title: str = "Heap Analytics Report",
        company_name: Optional[str] = None,
        logo_path: Optional[str] = None
    ):
        """
        Initialize Word document generator

        Args:
            output_path: Path for output DOCX file
            title: Report title
            company_name: Company name for header
            logo_path: Path to company logo image
        """
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for Word generation. Install with: pip install python-docx")

        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        self.title = title
        self.company_name = company_name
        self.logo_path = logo_path

        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """Configure document styles"""
        # Configure default styles
        style = self.doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(*self.COLORS["text"])

        # Title style
        style = self.doc.styles['Title']
        style.font.size = Pt(28)
        style.font.color.rgb = RGBColor(*self.COLORS["primary"])
        style.font.bold = True

        # Heading styles
        for i in range(1, 4):
            style = self.doc.styles[f'Heading {i}']
            style.font.color.rgb = RGBColor(*self.COLORS["primary"])

    def _set_cell_shading(self, cell, color_hex: str):
        """Set background color for a table cell"""
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), color_hex)
        cell._tc.get_or_add_tcPr().append(shading)

    def add_title_page(
        self,
        subtitle: Optional[str] = None,
        date_range: Optional[str] = None
    ):
        """Add title page to document"""
        # Logo
        if self.logo_path and os.path.exists(self.logo_path):
            self.doc.add_paragraph()
            logo_para = self.doc.add_paragraph()
            logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = logo_para.add_run()
            run.add_picture(self.logo_path, width=Inches(2))

        # Add spacing
        for _ in range(3):
            self.doc.add_paragraph()

        # Title
        title_para = self.doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(self.title)
        title_run.font.size = Pt(32)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(*self.COLORS["primary"])

        # Subtitle
        if subtitle:
            sub_para = self.doc.add_paragraph()
            sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = sub_para.add_run(subtitle)
            sub_run.font.size = Pt(16)
            sub_run.font.color.rgb = RGBColor(*self.COLORS["text_light"])

        self.doc.add_paragraph()

        # Company name
        if self.company_name:
            company_para = self.doc.add_paragraph()
            company_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            company_run = company_para.add_run(f"Prepared for: {self.company_name}")
            company_run.font.size = Pt(14)

        # Date range
        if date_range:
            date_para = self.doc.add_paragraph()
            date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            date_run = date_para.add_run(f"Reporting Period: {date_range}")
            date_run.font.size = Pt(12)
            date_run.font.color.rgb = RGBColor(*self.COLORS["text_light"])

        # Generation date
        gen_para = self.doc.add_paragraph()
        gen_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        gen_date = datetime.now().strftime("%B %d, %Y")
        gen_run = gen_para.add_run(f"Generated: {gen_date}")
        gen_run.font.size = Pt(12)
        gen_run.font.color.rgb = RGBColor(*self.COLORS["text_light"])

        # Page break
        self.doc.add_page_break()

    def add_section(self, title: str, level: int = 1):
        """Add a section heading"""
        heading = self.doc.add_heading(title, level=level)
        heading.style.font.color.rgb = RGBColor(*self.COLORS["primary"])

    def add_paragraph(self, text: str, bold: bool = False, italic: bool = False):
        """Add a paragraph of text"""
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.bold = bold
        run.font.italic = italic

    def add_bullet_list(self, items: List[str]):
        """Add a bulleted list"""
        for item in items:
            self.doc.add_paragraph(item, style='List Bullet')

    def add_page_break(self):
        """Add a page break"""
        self.doc.add_page_break()

    def add_executive_summary(self, summary: Dict[str, Any]):
        """Add executive summary section"""
        self.add_section("Executive Summary")

        # Key metrics table
        key_metrics = summary.get("key_metrics", {})
        if key_metrics:
            self._add_metrics_table(key_metrics)

        self.doc.add_paragraph()

        # Highlights
        highlights = summary.get("highlights", [])
        if highlights:
            self.add_section("Key Highlights", level=2)
            self.add_bullet_list(highlights)

    def _add_metrics_table(self, metrics: Dict[str, Any]):
        """Add a table of key metrics"""
        items = list(metrics.items())
        cols = min(4, len(items))
        rows = (len(items) + cols - 1) // cols

        table = self.doc.add_table(rows=rows * 2, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        idx = 0
        for row_idx in range(0, rows * 2, 2):
            for col_idx in range(cols):
                if idx < len(items):
                    name, value = items[idx]

                    # Value cell
                    value_cell = table.cell(row_idx, col_idx)
                    if isinstance(value, float):
                        display = f"{value:.1f}%" if "rate" in name.lower() else f"{value:,.1f}"
                    elif isinstance(value, int):
                        display = f"{value:,}"
                    else:
                        display = str(value)

                    value_para = value_cell.paragraphs[0]
                    value_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    value_run = value_para.add_run(display)
                    value_run.font.size = Pt(20)
                    value_run.font.bold = True
                    value_run.font.color.rgb = RGBColor(*self.COLORS["primary"])

                    # Label cell
                    label_cell = table.cell(row_idx + 1, col_idx)
                    label_para = label_cell.paragraphs[0]
                    label_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    label_run = label_para.add_run(name.replace("_", " ").title())
                    label_run.font.size = Pt(10)
                    label_run.font.color.rgb = RGBColor(*self.COLORS["text_light"])

                    idx += 1

    def add_chart(self, chart_path: str, title: Optional[str] = None, width: float = 6):
        """Add a chart image to the document"""
        if not os.path.exists(chart_path):
            return

        if title:
            self.add_section(title, level=2)

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(chart_path, width=Inches(width))

        self.doc.add_paragraph()

    def add_data_table(
        self,
        headers: List[str],
        data: List[List[Any]],
        title: Optional[str] = None
    ):
        """Add a data table"""
        if title:
            self.add_section(title, level=2)

        # Create table
        table = self.doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.style = 'Table Grid'

        # Header row
        header_row = table.rows[0]
        for i, header in enumerate(headers):
            cell = header_row.cells[i]
            cell.text = header
            # Style header
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.runs[0]
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            self._set_cell_shading(cell, '4F46E5')

        # Data rows
        for row_idx, row_data in enumerate(data):
            row = table.rows[row_idx + 1]
            for col_idx, value in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = str(value) if value is not None else ""

                # Alternate row colors
                if row_idx % 2 == 1:
                    self._set_cell_shading(cell, 'F9FAFB')

        self.doc.add_paragraph()

    def add_insights_section(self, insights: List[Dict[str, str]]):
        """Add insights section"""
        self.add_section("Key Insights")

        for insight in insights:
            title = insight.get("title", "")
            description = insight.get("description", "")
            priority = insight.get("priority", "medium")

            # Priority color
            color = {
                "high": self.COLORS["danger"],
                "medium": self.COLORS["accent"],
                "low": self.COLORS["secondary"]
            }.get(priority, self.COLORS["text"])

            # Title
            title_para = self.doc.add_paragraph()
            title_run = title_para.add_run(f"● {title}")
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(*color)

            # Description
            desc_para = self.doc.add_paragraph()
            desc_para.paragraph_format.left_indent = Inches(0.5)
            desc_para.add_run(description)

    def add_recommendations_section(self, recommendations: List[Dict[str, str]]):
        """Add recommendations section"""
        self.add_section("Recommendations")

        for i, rec in enumerate(recommendations, 1):
            area = rec.get("area", "General")
            priority = rec.get("priority", "Medium")
            text = rec.get("recommendation", "")
            impact = rec.get("impact", "")

            # Header
            header_para = self.doc.add_paragraph()
            header_run = header_para.add_run(f"{i}. {area} ")
            header_run.font.bold = True
            priority_run = header_para.add_run(f"(Priority: {priority})")
            priority_run.font.color.rgb = RGBColor(*self.COLORS["text_light"])

            # Recommendation text
            self.add_paragraph(text)

            # Impact
            if impact:
                impact_para = self.doc.add_paragraph()
                impact_para.paragraph_format.left_indent = Inches(0.25)
                impact_run = impact_para.add_run(f"Expected Impact: {impact}")
                impact_run.font.italic = True
                impact_run.font.color.rgb = RGBColor(*self.COLORS["secondary"])

            self.doc.add_paragraph()

    def add_page_analysis(self, page_analysis: Dict[str, Any]):
        """Add page performance analysis section"""
        self.add_section("Page Performance Analysis")

        total_pages = page_analysis.get("total_pages", 0)
        total_views = page_analysis.get("total_pageviews", 0)

        self.add_paragraph(
            f"Analyzed {total_pages} pages with {total_views:,} total pageviews.",
            bold=True
        )

        # Top pages table
        top_pages = page_analysis.get("top_pages_by_views", [])[:10]
        if top_pages:
            headers = ["Page URL", "Views", "Unique Visitors", "Avg Time"]
            data = []
            for page in top_pages:
                data.append([
                    page.get("url", "")[:40],
                    f"{page.get('total_views', 0):,}",
                    f"{page.get('unique_visitors', 0):,}",
                    f"{page.get('avg_time_on_page', 0):.0f}s"
                ])
            self.add_data_table(headers, data, "Top 10 Pages by Views")

        # High bounce pages
        high_bounce = page_analysis.get("pages_with_high_bounce", [])
        if high_bounce:
            self.add_section("Pages Needing Attention", level=2)
            self.add_paragraph(
                "The following pages have bounce rates above 70% and may need optimization:"
            )
            items = [
                f"{page.get('url', 'Unknown')}: {page.get('bounce_rate', 0):.1f}% bounce rate"
                for page in high_bounce[:5]
            ]
            self.add_bullet_list(items)

    def add_event_analysis(self, event_analysis: Dict[str, Any]):
        """Add event tracking analysis section"""
        self.add_section("Event Tracking Analysis")

        total_events = event_analysis.get("total_events_tracked", 0)
        total_count = event_analysis.get("total_event_count", 0)

        self.add_paragraph(
            f"Tracking {total_events} unique events with {total_count:,} total occurrences.",
            bold=True
        )

        # Top events table
        top_events = event_analysis.get("top_events", [])[:15]
        if top_events:
            headers = ["Event Name", "Type", "Count", "Unique Users"]
            data = []
            for event in top_events:
                data.append([
                    event.get("name", "")[:35],
                    event.get("event_type", "custom"),
                    f"{event.get('total_count', 0):,}",
                    f"{event.get('unique_users', 0):,}"
                ])
            self.add_data_table(headers, data, "Top Events")

    def add_journey_analysis(self, journey_analysis: Dict[str, Any]):
        """Add user journey analysis section"""
        self.add_section("User Journey Analysis")

        total_journeys = journey_analysis.get("total_journeys", 0)
        avg_length = journey_analysis.get("avg_journey_length", 0)

        self.add_paragraph(
            f"Analyzed {total_journeys} user journeys with an average of "
            f"{avg_length:.1f} steps per journey.",
            bold=True
        )

        # Entry points
        entry_points = journey_analysis.get("common_entry_points", [])
        if entry_points:
            self.add_section("Most Common Entry Points", level=2)
            items = [f"{entry}: {count:,} users" for entry, count in entry_points[:5]]
            self.add_bullet_list(items)

        # Exit points
        exit_points = journey_analysis.get("common_exit_points", [])
        if exit_points:
            self.add_section("Most Common Exit Points", level=2)
            items = [f"{exit_pt}: {count:,} users" for exit_pt, count in exit_points[:5]]
            self.add_bullet_list(items)

    def add_retention_analysis(self, retention_analysis: Dict[str, Any]):
        """Add retention analysis section"""
        self.add_section("User Retention Analysis")

        week1 = retention_analysis.get("week_1_retention", 0)
        health = retention_analysis.get("retention_health", "unknown")

        self.add_paragraph(
            f"Week 1 retention rate: {week1:.1f}% (Health status: {health.title()})",
            bold=True
        )

        # Recommendations
        recs = retention_analysis.get("recommendations", [])
        if recs:
            self.add_section("Retention Improvement Recommendations", level=2)
            self.add_bullet_list(recs)

    def generate(self) -> str:
        """
        Generate the Word document

        Returns:
            Path to generated DOCX file
        """
        self.doc.save(str(self.output_path))
        return str(self.output_path)

    def generate_full_report(
        self,
        analysis_results: Dict[str, Any],
        chart_paths: Dict[str, str],
        date_range_days: int = 30
    ) -> str:
        """
        Generate a complete report from analysis results

        Args:
            analysis_results: Results from AnalyticsProcessor.analyze_all()
            chart_paths: Dict of chart names to file paths
            date_range_days: Number of days in the report period

        Returns:
            Path to generated DOCX file
        """
        # Title page
        self.add_title_page(
            subtitle="Comprehensive Analytics Report",
            date_range=f"Last {date_range_days} days"
        )

        # Executive summary
        summary = analysis_results.get("summary", {})
        if summary:
            self.add_executive_summary(summary)

        # Charts section
        self.add_page_break()
        self.add_section("Visual Analytics")

        for chart_name, chart_path in chart_paths.items():
            if os.path.exists(chart_path):
                title = chart_name.replace("_", " ").title()
                self.add_chart(chart_path, title)

        # Page analysis
        self.add_page_break()
        page_analysis = analysis_results.get("page_analysis", {})
        if page_analysis and page_analysis.get("status") != "no_data":
            self.add_page_analysis(page_analysis)

        # Event analysis
        self.add_page_break()
        event_analysis = analysis_results.get("event_analysis", {})
        if event_analysis and event_analysis.get("status") != "no_data":
            self.add_event_analysis(event_analysis)

        # Journey analysis
        journey_analysis = analysis_results.get("journey_analysis", {})
        if journey_analysis and journey_analysis.get("status") != "no_data":
            self.add_journey_analysis(journey_analysis)

        # Retention analysis
        self.add_page_break()
        retention_analysis = analysis_results.get("retention_analysis", {})
        if retention_analysis and retention_analysis.get("status") != "no_data":
            self.add_retention_analysis(retention_analysis)

        # Insights
        insights = analysis_results.get("insights", [])
        if insights:
            self.add_page_break()
            self.add_insights_section(insights)

        # Recommendations
        recommendations = analysis_results.get("recommendations", [])
        if recommendations:
            self.add_page_break()
            self.add_recommendations_section(recommendations)

        return self.generate()
